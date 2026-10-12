"""
标准数据导入脚本
将 /doc/standard/ 目录下的 5 个 xlsx 文件导入到 SQLite 数据库中
使用 FTS5 全文搜索扩展建立索引
"""

import datetime
import os
import re
import sqlite3
from pathlib import Path

import openpyxl

# 项目路径
PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
DOC_DIR = PROJECT_ROOT / "doc" / "standard"
DB_DIR = Path(__file__).resolve().parent.parent / "database" / "standard"
DB_PATH = DB_DIR / "standards.db"

# 标准类别映射
CATEGORY_MAP = {
    "强制性国家标准.xlsx": "国标强制",
    "推荐性国家标准.xlsx": "国标推荐",
    "行业标准.xlsx": "行业标准",
    "地方标准.xlsx": "地方标准",
    "指导性技术文件.xlsx": "指导性技术文件",
}


def normalize_standard_no(raw: str) -> str:
    """
    标准号规范化处理
    1. 全角字符 → 半角字符
    2. 全角破折号/长横线 → 半角连字符
    3. 多余空格归一化
    4. 删除书名号
    5. 统一大写
    """
    if not raw:
        return ""

    s = raw.strip()

    # 全角字母和数字 → 半角
    result = []
    for ch in s:
        code = ord(ch)
        # 全角字母/数字（0xFF01-0xFF5E → 0x0021-0x007E）
        if 0xFF01 <= code <= 0xFF5E:
            result.append(chr(code - 0xFEE0))
        # 全角空格
        elif code == 0x3000:
            result.append(" ")
        else:
            result.append(ch)
    s = "".join(result)

    # 全角/特殊破折号 → 半角连字符
    s = re.sub(r"[—–―‐‑‒⁃]", "-", s)

    # 删除书名号
    s = s.replace("《", "").replace("》", "")
    s = s.replace("〈", "").replace("〉", "")

    # 多余空格归一化
    s = re.sub(r"\s+", " ", s).strip()

    # 统一大写
    s = s.upper()

    return s


def format_date(val) -> str:
    """格式化日期值为统一的字符串格式 YYYY-MM-DD"""
    if val is None:
        return ""
    if isinstance(val, datetime.datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, datetime.date):
        return val.strftime("%Y-%m-%d")
    # 字符串格式的日期，直接返回
    s = str(val).strip()
    if not s:
        return ""
    return s


def create_database(conn: sqlite3.Connection):
    """创建数据库表结构"""
    cursor = conn.cursor()

    # 主表
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS standards (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            standard_no TEXT NOT NULL,
            standard_no_normalized TEXT NOT NULL,
            name TEXT NOT NULL,
            category TEXT NOT NULL,
            is_adopted INTEGER DEFAULT 0,
            status TEXT DEFAULT '',
            publish_date TEXT DEFAULT '',
            implement_date TEXT DEFAULT '',
            industry TEXT DEFAULT '',
            region TEXT DEFAULT ''
        )
    """)

    # 创建普通索引
    cursor.execute(
        "CREATE INDEX IF NOT EXISTS idx_standard_no ON standards(standard_no)"
    )
    cursor.execute(
        "CREATE INDEX IF NOT EXISTS idx_standard_no_normalized ON standards(standard_no_normalized)"
    )
    cursor.execute(
        "CREATE INDEX IF NOT EXISTS idx_category ON standards(category)"
    )
    cursor.execute(
        "CREATE INDEX IF NOT EXISTS idx_status ON standards(status)"
    )

    # FTS5 全文搜索虚拟表
    cursor.execute("""
        CREATE VIRTUAL TABLE IF NOT EXISTS standards_fts USING fts5(
            standard_no_normalized,
            name,
            content='standards',
            content_rowid='id',
            tokenize='unicode61'
        )
    """)

    # FTS5 自动同步触发器
    cursor.execute("""
        CREATE TRIGGER IF NOT EXISTS standards_ai AFTER INSERT ON standards BEGIN
            INSERT INTO standards_fts(rowid, standard_no_normalized, name)
            VALUES (new.id, new.standard_no_normalized, new.name);
        END
    """)
    cursor.execute("""
        CREATE TRIGGER IF NOT EXISTS standards_ad AFTER DELETE ON standards BEGIN
            INSERT INTO standards_fts(standards_fts, rowid, standard_no_normalized, name)
            VALUES('delete', old.id, old.standard_no_normalized, old.name);
        END
    """)
    cursor.execute("""
        CREATE TRIGGER IF NOT EXISTS standards_au AFTER UPDATE ON standards BEGIN
            INSERT INTO standards_fts(standards_fts, rowid, standard_no_normalized, name)
            VALUES('delete', old.id, old.standard_no_normalized, old.name);
            INSERT INTO standards_fts(rowid, standard_no_normalized, name)
            VALUES (new.id, new.standard_no_normalized, new.name);
        END
    """)

    conn.commit()


def import_national_mandatory(conn: sqlite3.Connection, filepath: str):
    """导入强制性国家标准"""
    # 列：序号, 标准号, 是否采标, 标准名称, 状态, 发布日期, 实施日期, 操作
    _import_national(conn, filepath, "国标强制")


def import_national_recommended(conn: sqlite3.Connection, filepath: str):
    """导入推荐性国家标准"""
    # 列：序号, 标准号, 是否采标, 标准名称, 状态, 发布日期, 实施日期, 操作
    _import_national(conn, filepath, "国标推荐")


def _import_national(conn: sqlite3.Connection, filepath: str, category: str):
    """导入国家标准（强制性/推荐性），列结构相同"""
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    cursor = conn.cursor()
    count = 0
    batch = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:  # 跳过空行
            continue

        standard_no = str(row[1]).strip()
        is_adopted = 1 if row[2] and str(row[2]).strip() == "采" else 0
        name = str(row[3]).strip() if row[3] else ""
        status = str(row[4]).strip() if row[4] else ""
        publish_date = format_date(row[5])
        implement_date = format_date(row[6]) if len(row) > 6 else ""

        normalized = normalize_standard_no(standard_no)

        batch.append((
            standard_no, normalized, name, category,
            is_adopted, status, publish_date, implement_date, "", ""
        ))
        count += 1

        if len(batch) >= 1000:
            cursor.executemany(
                """INSERT INTO standards
                   (standard_no, standard_no_normalized, name, category,
                    is_adopted, status, publish_date, implement_date, industry, region)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                batch,
            )
            batch = []

    if batch:
        cursor.executemany(
            """INSERT INTO standards
               (standard_no, standard_no_normalized, name, category,
                is_adopted, status, publish_date, implement_date, industry, region)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            batch,
        )

    conn.commit()
    wb.close()
    print(f"  ✅ {category}: 导入 {count} 条记录")


def import_industry(conn: sqlite3.Connection, filepath: str):
    """导入行业标准"""
    # 列：#, 标准号, 标准名称, 行业领域, 状态, 批准日期, 实施日期
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    cursor = conn.cursor()
    count = 0
    batch = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue

        standard_no = str(row[1]).strip()
        name = str(row[2]).strip() if row[2] else ""
        industry = str(row[3]).strip() if row[3] else ""
        status = str(row[4]).strip() if row[4] else ""
        publish_date = format_date(row[5])
        implement_date = format_date(row[6]) if len(row) > 6 else ""

        normalized = normalize_standard_no(standard_no)

        batch.append((
            standard_no, normalized, name, "行业标准",
            0, status, publish_date, implement_date, industry, ""
        ))
        count += 1

        if len(batch) >= 1000:
            cursor.executemany(
                """INSERT INTO standards
                   (standard_no, standard_no_normalized, name, category,
                    is_adopted, status, publish_date, implement_date, industry, region)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                batch,
            )
            batch = []

    if batch:
        cursor.executemany(
            """INSERT INTO standards
               (standard_no, standard_no_normalized, name, category,
                is_adopted, status, publish_date, implement_date, industry, region)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            batch,
        )

    conn.commit()
    wb.close()
    print(f"  ✅ 行业标准: 导入 {count} 条记录")


def import_local(conn: sqlite3.Connection, filepath: str):
    """导入地方标准"""
    # 列：#, 标准号, 标准名称, 省市区, 状态, 批准日期, 实施日期
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    cursor = conn.cursor()
    count = 0
    batch = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue

        standard_no = str(row[1]).strip()
        name = str(row[2]).strip() if row[2] else ""
        region = str(row[3]).strip() if row[3] else ""
        status = str(row[4]).strip() if row[4] else ""
        publish_date = format_date(row[5])
        implement_date = format_date(row[6]) if len(row) > 6 else ""

        normalized = normalize_standard_no(standard_no)

        batch.append((
            standard_no, normalized, name, "地方标准",
            0, status, publish_date, implement_date, "", region
        ))
        count += 1

        if len(batch) >= 1000:
            cursor.executemany(
                """INSERT INTO standards
                   (standard_no, standard_no_normalized, name, category,
                    is_adopted, status, publish_date, implement_date, industry, region)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                batch,
            )
            batch = []

    if batch:
        cursor.executemany(
            """INSERT INTO standards
               (standard_no, standard_no_normalized, name, category,
                is_adopted, status, publish_date, implement_date, industry, region)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            batch,
        )

    conn.commit()
    wb.close()
    print(f"  ✅ 地方标准: 导入 {count} 条记录")


def import_guidelines(conn: sqlite3.Connection, filepath: str):
    """导入指导性技术文件"""
    # 列：序号, 标准号, 是否采标, 标准名称, 状态, 发布日期, 操作
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    cursor = conn.cursor()
    count = 0
    batch = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue

        standard_no = str(row[1]).strip()
        is_adopted = 1 if row[2] and str(row[2]).strip() == "采" else 0
        name = str(row[3]).strip() if row[3] else ""
        status = str(row[4]).strip() if row[4] else ""
        publish_date = format_date(row[5])

        normalized = normalize_standard_no(standard_no)

        batch.append((
            standard_no, normalized, name, "指导性技术文件",
            is_adopted, status, publish_date, "", "", ""
        ))
        count += 1

        if len(batch) >= 1000:
            cursor.executemany(
                """INSERT INTO standards
                   (standard_no, standard_no_normalized, name, category,
                    is_adopted, status, publish_date, implement_date, industry, region)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                batch,
            )
            batch = []

    if batch:
        cursor.executemany(
            """INSERT INTO standards
               (standard_no, standard_no_normalized, name, category,
                is_adopted, status, publish_date, implement_date, industry, region)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            batch,
        )

    conn.commit()
    wb.close()
    print(f"  ✅ 指导性技术文件: 导入 {count} 条记录")


def print_stats(conn: sqlite3.Connection):
    """打印导入统计"""
    cursor = conn.cursor()

    cursor.execute("SELECT COUNT(*) FROM standards")
    total = cursor.fetchone()[0]
    print("\n📊 数据库统计:")
    print(f"  总记录数: {total}")

    cursor.execute(
        "SELECT category, COUNT(*) FROM standards GROUP BY category ORDER BY COUNT(*) DESC"
    )
    for row in cursor.fetchall():
        print(f"  {row[0]}: {row[1]}")

    cursor.execute(
        "SELECT status, COUNT(*) FROM standards GROUP BY status ORDER BY COUNT(*) DESC"
    )
    print("\n  状态分布:")
    for row in cursor.fetchall():
        print(f"    {row[0]}: {row[1]}")


def main():
    """主函数"""
    print("🚀 开始导入标准数据...")
    print(f"  数据目录: {DOC_DIR}")
    print(f"  数据库路径: {DB_PATH}")

    # 确保目录存在
    DB_DIR.mkdir(parents=True, exist_ok=True)

    # 删除旧数据库
    if DB_PATH.exists():
        os.remove(DB_PATH)
        print("  🗑️ 已删除旧数据库")

    # 创建数据库连接
    conn = sqlite3.connect(str(DB_PATH))
    # 性能优化
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    conn.execute("PRAGMA cache_size=-64000")  # 64MB 缓存

    try:
        # 创建表结构
        create_database(conn)
        print("  📋 数据库表结构已创建")

        # 导入数据
        print("\n📥 导入数据:")
        import_national_mandatory(conn, str(DOC_DIR / "强制性国家标准.xlsx"))
        import_national_recommended(conn, str(DOC_DIR / "推荐性国家标准.xlsx"))
        import_industry(conn, str(DOC_DIR / "行业标准.xlsx"))
        import_local(conn, str(DOC_DIR / "地方标准.xlsx"))
        import_guidelines(conn, str(DOC_DIR / "指导性技术文件.xlsx"))

        # 打印统计
        print_stats(conn)

        print("\n✅ 数据导入完成!")

    finally:
        conn.close()


if __name__ == "__main__":
    main()
