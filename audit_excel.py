"""Export audit results to filterable Excel sheets; retain full records in JSONL."""
import csv
import datetime as dt
import json
import math
import re
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from standard_audit import SOURCES, FIELD_LABELS


def supplement_workbook(input_path, result_dir, output_path):
    """Append audit columns to a copy, preserving worksheets, values, formulas and styles."""
    from standard_audit import normalize

    result_dir, output_path = Path(result_dir), Path(output_path)
    with (result_dir / "audit-results.csv").open(encoding="utf-8-sig", newline="") as f:
        results = list(csv.DictReader(f))
    book = load_workbook(input_path)
    try:
        sheet = next((s for s in book if "标准号" in [c.value for c in s[1]]), None)
        if sheet is None or sheet.max_row != len(results) + 1:
            raise ValueError("Audit row count differs from the original worksheet; refusing to merge")
        headers = {c.value: c.column for c in sheet[1] if c.value}
        fields = ["Verified standard name", "Verified status", "Verified publication date", "Verified implementation date", "Verified withdrawal date",
                  "Verified current versions", "Current version publication dates", "Current version implementation dates",
                  "Source URLs", "Checked at", "Conclusion", "Audit notes"]
        for name in fields:
            if name in headers:
                raise ValueError("Use the original workbook without appended audit columns")
            column = sheet.max_column + 1
            cell = sheet.cell(1, column, name)
            cell._style = copy(sheet.cell(1, headers["标准号"])._style)
            sheet.column_dimensions[get_column_letter(column)].width = 42 if name in fields[-4:] else 24
            headers[name] = column
        for number, result in enumerate(results, 2):
            if int(result["Input row"]) != number or normalize(result["标准号"]) != normalize(sheet.cell(number, headers["标准号"]).value):
                raise ValueError("Audit row or standard number does not match the original; refusing to merge")
            for name in fields:
                cell = sheet.cell(number, headers[name])
                text_cell(cell, result.get(name, ""))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                value = result.get(name, "")
                if name.endswith(("日期", "date")) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
                    cell.value = dt.date.fromisoformat(value)
                    cell.number_format = "yyyy-mm-dd"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        temp = output_path.with_suffix(".tmp.xlsx")
        book.save(temp)
        temp.replace(output_path)
        return output_path
    finally:
        book.close()

DETAIL_COLUMNS = {"Field provenance": "Field evidence", "Current version provenance": "Field evidence", "Replacement chain": "Replacement relationships"}


def text_cell(cell, value):
    text = str(value or "")
    if len(text) > 32767:
        text = "Content exceeds the Excel cell limit; see audit-records.jsonl for the full record."
    cell.value = text
    cell.data_type = "s"


def style_table(sheet, header_row=1, freeze="A2"):
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for row in sheet.iter_rows(min_row=header_row):
        height = 1
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="23313D")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row[0].row == header_row:
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="24465B")
            elif row[0].row % 2:
                cell.fill = PatternFill("solid", fgColor="F2F6F8")
            width = sheet.column_dimensions[cell.column_letter].width
            height = max(height, sum(max(1, math.ceil(len(s) / max(8, width * .65))) for s in str(cell.value or "").split("\n")))
        sheet.row_dimensions[row[0].row].height = min(300, max(30, height * 14 + 6))
    sheet.sheet_view.showGridLines = False


def export_excel(directory):
    directory = Path(directory)
    with (directory / "audit-results.csv").open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        headers, rows = reader.fieldnames, list(reader)
    records = [json.loads(s) for s in (directory / "audit-records.jsonl").read_text().splitlines()]
    summary_path = directory / "run-summary.json"
    summary = json.loads(summary_path.read_text()) if summary_path.exists() else {}
    workbook = Workbook()
    result = workbook.active
    result.title = "Audit results"
    result.merge_cells("A1:H1")
    result["A1"] = "Standards audit results"
    result["A1"].font = Font(name="Arial", size=18, bold=True, color="24465B")
    result.row_dimensions[1].height = 32
    result.merge_cells("A2:H2")
    result["A2"] = summary.get("task_status", "Baseline screening / partial online audit") + "; unmatched and unsupported records are retained. Review the Conclusion column."
    result["A2"].font = Font(name="Arial", size=11, color="475A69")
    result["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    result.row_dimensions[2].height = 36
    for column, name in enumerate(headers, 1):
        result.cell(4, column, name)
        width = 18
        if "标准名称" in name or "standard name" in name or name == "文件名":
            width = 44
        elif "来源" in name or "说明" in name or name in ("备注", "Source URLs", "Audit notes", "Conclusion"):
            width = 62
        elif "现行版本" in name or "version" in name:
            width = 36
        elif name in DETAIL_COLUMNS:
            width = 24
        result.column_dimensions[get_column_letter(column)].width = width
    for number, row in enumerate(rows, 5):
        for column, name in enumerate(headers, 1):
            value = row[name]
            cell = result.cell(number, column)
            if name in DETAIL_COLUMNS and value:
                value = "See the " + DETAIL_COLUMNS[name] + " sheet and audit-records.jsonl"
            if name.endswith(("日期", "date")) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value or ""):
                try:
                    cell.value = dt.date.fromisoformat(value)
                    cell.number_format = "yyyy-mm-dd"
                except ValueError:
                    text_cell(cell, value)
            else:
                text_cell(cell, value)
    style_table(result, 4, "D5")
    conclusion_column = headers.index("Conclusion") + 1
    for number in range(5, result.max_row + 1):
        cell = result.cell(number, conclusion_column)
        color = {"Verified by official sources": "157A64", "Manual review required": "9D5D06", "Unconfirmed": "A03838"}.get(cell.value)
        if color:
            cell.font = Font(name="Arial", size=10, bold=True, color=color)
    # Flatten evidence by input row, field and source instead of embedding large JSON cells.
    evidence = workbook.create_sheet("Field evidence")
    evidence.append(["Input row", "Original standard number", "Audited standard number", "Field", "Value", "Source", "Selected", "Fetched at", "Source URL", "Snapshot path", "SHA256", "Request parameters"])
    relations = workbook.create_sheet("Replacement relationships")
    relations.append(["Input row", "Original standard number", "Standard number in chain", "Status", "Replaced by", "Conflicts / lookup notes", "Relation text and sources"])
    evidence_row = relation_row = 1
    for record in records:
        original = record["original"]
        row_number = original.get("Original worksheet row", record["Input row"])
        raw_code = original["标准号"]
        items = [(raw_code, record.get("provenance", {}))]
        items.extend((v["标准号"], v["Field provenance"]) for v in record.get("current_versions", []))
        for code, fields in items:
            for field, sources in fields.items():
                for source in sources:
                    values = [row_number, raw_code, code, FIELD_LABELS.get(field, field), source.get("value", ""), SOURCES[source["source"]][0],
                              "Yes" if source.get("selected") else "No / relation evidence", source["fetched_at"], source["url"],
                              source["snapshot"], source["sha256"], source.get("request", {})]
                    evidence.append([json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v for v in values])
                    evidence_row += 1
                    for cell in (evidence.cell(evidence_row, c) for c in range(1, 13)):
                        text_cell(cell, cell.value)
                    if source["url"].startswith(("https://", "http://")):
                        evidence.cell(evidence_row, 9).hyperlink = source["url"]
        for link in record.get("replacement_chain", []):
            rel_sources = link.get("Field provenance", {}).get("Replacement relationships", [])
            explanation = link.get("conflicts", []) + [e["message"] for e in link.get("errors", [])]
            relations.append([row_number, raw_code, link["标准号"], link["状态"], "；".join(link["替代标准"]),
                              "；".join(explanation), json.dumps(rel_sources, ensure_ascii=False)])
            relation_row += 1
            for cell in (relations.cell(relation_row, c) for c in range(1, 8)):
                text_cell(cell, cell.value)
    for sheet in (evidence, relations):
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 22 if column < 5 else 58
        style_table(sheet)
    target = directory / "audit-results.xlsx"
    temporary = target.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    workbook.close()
    # Check saved row counts and reject unexpected formulas.
    saved = load_workbook(temporary, read_only=True, data_only=False)
    try:
        sheet = saved["Audit results"]
        if sheet.max_row != len(rows) + 4:
            raise ValueError("Excel export row count mismatch")
        for values in saved:
            for row in values:
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("Unexpected formula in the exported audit workbook")
    finally:
        saved.close()
    temporary.replace(target)
    return target
