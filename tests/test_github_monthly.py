"""月度流程回归：失败页、分片恢复、原表保护与 GitHub 迁移。"""
import csv
import json
from argparse import Namespace
from pathlib import Path
from unittest.mock import Mock

import pytest
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

import standard_crawl as crawler
from audit_excel import supplement_workbook
from scripts import export_github, github_monthly
from standard_audit import Evidence, Observation, merge_observations, now


def test_only_selected_tasks_can_finish(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(crawler, "get_page_data_sacinfo", lambda *a: ([{"标准号": "NB/T 1-2020", "状态": "现行"}], 1))
    args = crawler.parse_args(["--only", "industry", "--strict-pages", "--delay", "0"])
    status = crawler.run_chunk(args)
    assert status["finished"] and status["next_action"] == "build"
    assert not status["tasks"]["local"]["finished"]


def test_failed_page_preserves_previous_page_and_rows(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    fetch = Mock(side_effect=[([{"标准号": "NB/T 1-2020"}], 2), (None, 0)])
    monkeypatch.setattr(crawler, "get_page_data_sacinfo", fetch)
    args = crawler.parse_args(["--only", "industry", "--strict-pages", "--page-retries", "1", "--delay", "0"])
    status = crawler.run_chunk(args)
    assert status["next_action"] == "stop"
    assert status["tasks"]["industry"]["page"] == 1
    assert status["tasks"]["industry"]["skipped_pages"] == []
    assert "NB/T 1-2020" in Path("行业公开标准信息.csv").read_text()
    monkeypatch.setattr(crawler, "get_page_data_sacinfo", lambda *a: ([{"标准号": "NB/T 2-2020"}], 2))
    assert crawler.run_chunk(args)["finished"]
    assert "NB/T 1-2020" in Path("行业公开标准信息.csv").read_text()


@pytest.mark.parametrize("payload", [{}, {"records": [], "pages": 3}, {"records": [{"chName": "无标准号"}], "pages": 1}, []])
def test_bad_api_response_is_not_treated_as_complete(payload, monkeypatch):
    response = Mock()
    response.json.return_value = payload
    monkeypatch.setattr(crawler, "fetch_with_retry", lambda *a, **k: response)
    assert crawler.get_page_data_sacinfo(crawler.TASK_BY_KEY["local"], 1, 50)[0] is None


def make_workbook(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "标准信息"
    sheet.append(["标准号", "状态", "备注"])
    sheet.append(["GB 1-2020", "原状态", "=1+1"])
    sheet["A2"].font = Font(bold=True, color="FF0000")
    book.create_sheet("说明")["A1"] = "保留说明"
    book.save(path)
    book.close()


def test_supplement_preserves_sheets_values_formulas_and_literal_text(tmp_path):
    source = tmp_path / "原表.xlsx"
    make_workbook(source)
    before = source.read_bytes()
    with (tmp_path / "标准信息核验结果.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=["标准号", "原表行号", "核验状态", "核验发布日期", "核验标准名称"])
        writer.writeheader()
        writer.writerow({"标准号": "GB 1-2020", "原表行号": 2, "核验状态": "废止",
                         "核验发布日期": "2020-01-01", "核验标准名称": "=外部文本"})
    out = supplement_workbook(source, tmp_path, tmp_path / "补充.xlsx")
    book = load_workbook(out)
    sheet = book["标准信息"]
    assert book.sheetnames == ["标准信息", "说明"]
    assert sheet["B2"].value == "原状态" and sheet["C2"].data_type == "f"
    assert sheet["A2"].font.bold
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, headers["核验状态"]).value == "废止"
    assert sheet.cell(2, headers["核验标准名称"]).data_type == "s"
    assert sheet.cell(2, headers["核验发布日期"]).number_format == "yyyy-mm-dd"
    assert source.read_bytes() == before
    book.close()


def test_audit_restore_requires_matching_month_and_input(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    make_workbook(github_monthly.INPUT)
    args = Namespace(campaign="2026-10", restore=True, budget_seconds=60, delay=0)
    with pytest.raises(ValueError, match="缺少前片"):
        github_monthly.audit(args)
    github_monthly.AUDIT_STATE.mkdir()
    meta = github_monthly.AUDIT_STATE / "campaign.json"
    meta.write_text(json.dumps({"campaign": "2026-09", "input_sha256": "old", "chunks": 1}))
    with pytest.raises(ValueError, match="月份或原表"):
        github_monthly.audit(args)


def test_audit_end_to_end_with_fresh_online_result(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    make_workbook(github_monthly.INPUT)
    with Path("国家公开标准信息.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=crawler.NATIONAL_COLUMNS)
        writer.writeheader()
        writer.writerow({"标准类型": "强制性国家标准", "标准号": "GB 1-2020", "标准名称": "测试标准",
                         "状态": "现行", "发布日期": "2020-01-01", "实施日期": "2020-02-01"})
    def online(code):
        ev = Evidence("samr", "https://std.samr.gov.cn/test", now(), "网页证据/test.txt", "test")
        data = merge_observations([Observation(code, {"状态": "废止", "标准名称": "测试标准"}, [], [], ev)], [])
        data.update({"current_versions": [], "replacement_chain": [], "review_notes": ["未确认替代版本"]})
        return data
    lookup = Mock(side_effect=online)
    monkeypatch.setattr(github_monthly.audit_pipeline.Auditor, "audit", lambda self, code: lookup(code))
    args = Namespace(campaign="2026-10", restore=False, budget_seconds=60, delay=0)
    github_monthly.audit(args)
    assert lookup.call_count == 1  # 历史库命中也必须联网。
    result = json.loads((github_monthly.AUDIT_STATE / "results/运行摘要.json").read_text())
    assert result["自动续跑"]  # 未确认项允许第二次尝试。
    args.restore = True
    github_monthly.audit(args)
    assert lookup.call_count == 2
    assert Path("reports/标准信息统计表_已核验.xlsx").exists()
    assert Path(".github_release/核验证据.tar.gz").exists()
    book = load_workbook("reports/标准信息统计表_已核验.xlsx")
    assert book["标准信息"]["B2"].value == "原状态"
    headers = {c.value: c.column for c in book["标准信息"][1]}
    assert book["标准信息"].cell(2, headers["核验状态"]).value == "废止"
    book.close()


def test_export_excludes_history_database_and_local_state(tmp_path, monkeypatch):
    root = tmp_path / "source"
    root.mkdir()
    names = ["main.py", "标准信息统计表.xlsx", "audit_inputs/manifest.json", "standard_db/standards.db", ".git/config", "local_audit/private.txt"]
    for name in names:
        path = root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("样本")
    monkeypatch.setattr(export_github, "ROOT", root)
    monkeypatch.setattr(export_github.subprocess, "check_output", lambda *a, **k: "\0".join(names).encode())
    destination = tmp_path / "destination"
    export_github.export(destination)
    assert sorted(p.name for p in destination.iterdir()) == ["audit_inputs", "main.py", "标准信息统计表.xlsx"]
    assert (destination / "audit_inputs/manifest.json").exists()
    with pytest.raises(ValueError, match="已存在"):
        export_github.export(destination)


def test_workflows_have_bounded_jobs_and_no_push_crawl_trigger():
    root = Path(__file__).resolve().parents[1]
    for name in ("monthly-audit", "monthly-crawl"):
        # BaseLoader 保留 on 的字符串语义，避免 YAML 1.1 把 on 当作布尔值。
        config = yaml.load((root / f".github/workflows/{name}.yml").read_text(), Loader=yaml.BaseLoader)
        assert set(config["on"]) == {"schedule", "workflow_dispatch"}
        assert config["permissions"] == {"contents": "write", "actions": "write"}
        job = next(iter(config["jobs"].values()))
        assert int(job["timeout-minutes"]) < 360
        assert job["if"] == "github.ref == 'refs/heads/main'"
        assert "github.token" in job["env"]["GH_TOKEN"]
