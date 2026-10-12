"""公开清单不携带私有字段，本地回填保留原表信息。"""
import csv

import pytest
from openpyxl import Workbook, load_workbook

from scripts.public_input import merge_local, prepare_public, validate_public


def original(path):
    book = Workbook()
    book.properties.creator = "不应公开的作者"
    sheet = book.active
    sheet.title = "原表"
    sheet.append(["标准类型", "标准号", "文件名", "备注"])
    sheet.append(["国家标准", "GB 2-2020", "私有甲.pdf", "内部备注"])
    sheet.append(["国家标准", "GB 1-2020", "私有乙.pdf", "=1+1"])
    sheet.append(["国家标准", "GB 2-2020", "私有丙.pdf", "内部备注"])
    book.save(path)
    book.close()


def test_public_input_drops_metadata_private_columns_and_row_order(tmp_path):
    source, target = tmp_path / "original.xlsx", tmp_path / "public.csv"
    original(source)
    before = source.read_bytes()
    assert prepare_public(source, target) == 2
    assert validate_public(target) == [
        {"标准类型": "国内标准", "标准号": "GB1-2020"},
        {"标准类型": "国内标准", "标准号": "GB2-2020"},
    ]
    assert source.read_bytes() == before
    for text in ("作者", "文件名", "备注", "行号", ".pdf"):
        assert text not in target.read_text(encoding="utf-8-sig")


def test_rejects_private_columns_and_formula_input(tmp_path):
    path = tmp_path / "public.csv"
    path.write_text("标准类型,标准号,备注\n国内标准,GB1-2020,隐私\n")
    with pytest.raises(ValueError, match="只能包含"):
        validate_public(path)
    path.write_text("标准类型,标准号\n国内标准,=1+1\n")
    with pytest.raises(ValueError, match="非法"):
        validate_public(path)


def test_local_merge_preserves_private_fields_duplicates_and_missing_results(tmp_path):
    source = tmp_path / "original.xlsx"
    original(source)
    results = tmp_path / "results.csv"
    with results.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["标准类型", "标准号", "原表行号", "核验状态", "核验结论"])
        writer.writeheader()
        writer.writerow({"标准类型": "国内标准", "标准号": "GB2-2020", "原表行号": 2, "核验状态": "现行", "核验结论": "官方来源已核验"})
    output = merge_local(source, results, tmp_path / "local.xlsx")
    book = load_workbook(output)
    sheet = book["原表"]
    headers = {c.value: c.column for c in sheet[1]}
    assert sheet["C2"].value == "私有甲.pdf" and sheet["C4"].value == "私有丙.pdf"
    assert sheet["D3"].data_type == "f"
    assert sheet.cell(2, headers["核验状态"]).value == sheet.cell(4, headers["核验状态"]).value == "现行"
    assert sheet.cell(3, headers["核验结论"]).value == "未采集"
    book.close()
    with pytest.raises(ValueError, match="不能覆盖"):
        merge_local(source, results, source)
