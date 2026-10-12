"""自动续跑、跨提交恢复、Excel 日期及文本安全的离线验证。"""
import json
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from audit_excel import export_excel
from audit_pipeline import parse_args, precheck, prepare, run
from standard_audit import Fetcher, LocalStore
from test_audit_pipeline import make_fixture


class DeliveryTests(unittest.TestCase):


    def test_excel_keeps_dates_duplicate_rows_evidence_and_literal_text(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            db, input_path = make_fixture(root)
            input_path.write_text(input_path.read_text(encoding="utf-8-sig").replace("甲.pdf", "=1+1"), encoding="utf-8-sig")
            precheck(input_path, db, root / "out")
            path = export_excel(root / "out")
            book = load_workbook(path)
            try:
                self.assertEqual(book.sheetnames, ["标准核验结果", "字段证据", "替代关系"])
                sheet = book.worksheets[0]
                self.assertEqual(sheet.max_row, 9)
                self.assertEqual(sheet["A6"].value, sheet["A7"].value)
                self.assertEqual(sheet["D6"].data_type, "s")
                self.assertIn("=1+1", sheet["D6"].value)
                self.assertEqual(sheet["C5"].number_format, "yyyy-mm-dd")
                self.assertEqual(sheet.freeze_panes, "D5")
                self.assertGreater(book["字段证据"].max_row, 1)
                self.assertFalse(any(cell.data_type == "f" for s in book for row in s for cell in row))
            finally:
                book.close()

    def test_site_pause_survives_a_new_fetcher(self):
        with tempfile.TemporaryDirectory() as directory:
            store = LocalStore(directory)
            try:
                first = Fetcher(store, delay=0)
                first.pause("openstd.samr.gov.cn")
                second = Fetcher(store, delay=0)
                second.session.request = Mock()
                with self.assertRaisesRegex(ValueError, "已暂停"):
                    second.fetch("national", "https://openstd.samr.gov.cn/test")
                second.session.request.assert_not_called()
            finally:
                store.conn.close()

    def test_no_progress_or_all_sources_paused_stops_automatic_chain(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            db, input_path = make_fixture(root)
            prepare(parse_args(["prepare", str(input_path), "--db", str(db), "--output-dir", str(root / "pre"),
                                "--manifest-dir", str(root / "inputs")]))
            output = root / "cloud"
            output.mkdir()
            hosts = ["std.samr.gov.cn", "openstd.samr.gov.cn", "hbba.sacinfo.org.cn", "dbba.sacinfo.org.cn"]
            (output / "站点暂停状态.json").write_text(json.dumps({h: time.time() + 3600 for h in hosts}))
            args = parse_args(["run", str(root / "inputs/standards.csv"), "--db", str(db),
                               "--manifest", str(root / "inputs/manifest.json"), "--output-dir", str(output)])
            with patch("audit_pipeline.Auditor") as factory:
                factory.return_value.audit.return_value = {"values": {}, "observations": [], "errors": [], "conflicts": [], "review_notes": []}
                self.assertEqual(run(args), 2)
            summary = json.loads((output / "运行摘要.json").read_text())
            self.assertFalse(summary["自动续跑"])
            self.assertIn("全部官方站点", summary["停止原因"])


if __name__ == "__main__":
    unittest.main()
